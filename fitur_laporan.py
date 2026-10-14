"""UNTUK APA: laporan rekap + unduh Excel + backup + hapus data.
- GET  /api/ringkas?dari&sampai&q&batas -> rekap per orang per hari
- GET  /api/demo? -> no. Demo ada di fitur_tarik (/api/demo).
- GET  /export?dari&sampai&tgl&batas&format=xlsx|csv -> unduh Excel/CSV
- GET  /backup -> unduh database (.db saat SQLite, .sql saat MySQL, khusus admin)
- POST /api/reset {scope, date, dry} -> hapus harian/mingguan/bulanan/tahunan (khusus admin,
  otomatis backup dulu). Kirim dry:true dulu untuk hitung jumlah.
Admin saja. User tidak boleh akses laporan (rekap maupun unduh)."""
import io
import os
from calendar import monthrange
from datetime import date, datetime, timedelta
from flask import Blueprint, request, jsonify, Response, session, send_file
from database import db, snapshot_db, pakai_mysql, dump_sql
from auth import wajib_admin
from config import DB, wib_now, wib_today, tentukan_shift

bp_laporan = Blueprint("laporan", __name__)
BACKUP_DIR = os.path.join(os.path.dirname(os.path.abspath(DB)), "backup")
LABEL = {"tepat": "Tepat waktu", "telat": "Terlambat", "satu": "Hanya 1 scan",
         "belum": "Belum hadir", "izin": "Izin/Cuti"}


def _tanggal(teks, cadangan):
    try:
        return date.fromisoformat(teks)
    except (TypeError, ValueError):
        return cadangan


def _rentang():
    dari = _tanggal(request.args.get("dari") or request.args.get("tgl") or request.args.get("date"),
                    wib_today())
    sampai = _tanggal(request.args.get("sampai"), dari)
    return (dari, sampai) if dari <= sampai else (sampai, dari)


def _hitung(dari, sampai, q="", batas="08:15", urut="tanggal"):
    """Satu baris per pegawai per hari: masuk = scan pertama, pulang = scan terakhir.
    urut='tanggal' -> tanggal 1-30 dulu baru nama (rapih per hari).
    urut='nama' -> nama abjad dulu baru tanggal."""
    if session.get("role") != "admin":
        sendiri = (session.get("nik") or "").strip()
        if not sendiri:
            return []  # NIK belum diikat: tidak boleh lihat apa-apa
        q = ""
    else:
        sendiri = ""
    like = f"%{q}%"
    c = db()
    # Hanya absen DISETUJUI yang masuk rekap (PENDING lokasi menunggu ACC,
    # DITOLAK tidak dihitung — sama seperti izin yang belum disetujui).
    rows = c.execute(
        """SELECT date(a.waktu) tgl, a.nik, COALESCE(p.nama,'') nama, COALESCE(p.bagian,'') bagian,
                   MIN(a.waktu) pertama, MAX(a.waktu) terakhir, COUNT(*) jml
           FROM absen a LEFT JOIN pegawai p ON p.nik = a.nik
           WHERE date(a.waktu) BETWEEN ? AND ?
             AND COALESCE(a.status,'DISETUJUI')='DISETUJUI'
             AND (? = '' OR a.nik LIKE ? OR COALESCE(p.nama,'') LIKE ? OR COALESCE(p.bagian,'') LIKE ?)
             AND (? = '' OR a.nik = ?)
           GROUP BY date(a.waktu), a.nik ORDER BY COALESCE(p.nama,a.nik), tgl""",
        (dari.isoformat(), sampai.isoformat(), q, like, like, like, sendiri, sendiri)).fetchall()
    daftar = c.execute("SELECT nik, nama, COALESCE(bagian,'') bagian FROM pegawai"
                       + (" WHERE nik=?" if sendiri else "")
                       + " ORDER BY COALESCE(nama,nik)",
                       ((sendiri,) if sendiri else ())).fetchall()
    izin = {}
    if sendiri:
        for r in c.execute("SELECT nik, tanggal, jenis FROM pengajuan "
                           "WHERE status='DISETUJUI' AND nik=? AND tanggal BETWEEN ? AND ?",
                           (sendiri, dari.isoformat(), sampai.isoformat())).fetchall():
            izin[(r["nik"], r["tanggal"])] = r["jenis"]
    elif not q:
        for r in c.execute("SELECT nik, tanggal, jenis FROM pengajuan "
                           "WHERE status='DISETUJUI' AND tanggal BETWEEN ? AND ?",
                           (dari.isoformat(), sampai.isoformat())).fetchall():
            izin[(r["nik"], r["tanggal"])] = r["jenis"]
    c.close()
    out = []
    for r in rows:
        masuk, pulang = r["pertama"][11:16], r["terakhir"][11:16]
        menit = None
        if r["jml"] > 1:
            f = "%Y-%m-%d %H:%M:%S"
            menit = int((datetime.strptime(r["terakhir"], f)
                         - datetime.strptime(r["pertama"], f)).total_seconds() // 60)
        if (r["nik"], r["tgl"]) in izin:
            st = "izin"
        else:
            st = "satu" if r["jml"] == 1 else ("telat" if masuk > batas else "tepat")
        out.append({"tanggal": r["tgl"], "nik": r["nik"], "nama": r["nama"], "bagian": r["bagian"],
                    "masuk": masuk, "pulang": pulang if r["jml"] > 1 else "",
                    "shift": tentukan_shift(masuk),
                    "jml": r["jml"], "menit": menit, "status": st})
    if dari == sampai and not q:  # pegawai terdaftar yang belum scan hari itu
        ada = {(r["nik"], r["tanggal"]) for r in
               [{"nik": x["nik"], "tanggal": x["tanggal"]} for x in out]}
        for p in daftar:
            if (p["nik"], dari.isoformat()) in ada:
                continue
            if (p["nik"], dari.isoformat()) in izin:
                out.append({"tanggal": dari.isoformat(), "nik": p["nik"], "nama": p["nama"],
                            "bagian": p["bagian"], "masuk": "", "pulang": "",
                            "shift": "",
                            "jml": 0, "menit": None, "status": "izin"})
            else:
                out.append({"tanggal": dari.isoformat(), "nik": p["nik"], "nama": p["nama"],
                            "bagian": p["bagian"], "masuk": "", "pulang": "",
                            "shift": "",
                            "jml": 0, "menit": None, "status": "belum"})
    # Urut: tanggal 1-30 dulu baru nama (rapih per hari), atau nama dulu.
    if urut == "nama":
        out.sort(key=lambda x: ((x["nama"] or x["nik"] or "").upper(), x["tanggal"]))
    else:
        out.sort(key=lambda x: (x["tanggal"], (x["nama"] or x["nik"] or "").upper()))
    return out


def _durasi(menit):
    return "" if menit is None else f"{menit // 60}j {menit % 60:02d}m"


@bp_laporan.route("/api/ringkas")
@wajib_admin
def api_ringkas():
    dari, sampai = _rentang()
    urut = (request.args.get("urut") or "tanggal").strip().lower()
    if urut not in ("tanggal", "nama"):
        urut = "tanggal"
    return jsonify(_hitung(dari, sampai, (request.args.get("q") or "").strip(),
                           request.args.get("batas", "08:15") or "08:15", urut))


@bp_laporan.route("/export")
@wajib_admin
def export():
    dari, sampai = _rentang()
    batas = request.args.get("batas", "08:15") or "08:15"
    urut = (request.args.get("urut") or "tanggal").strip().lower()
    if urut not in ("tanggal", "nama"):
        urut = "tanggal"
    q = (request.args.get("q") or "").strip()
    if session.get("role") != "admin":
        q = ""
    data = _hitung(dari, sampai, q, batas, urut)
    if (request.args.get("format") or "xlsx").lower() == "csv":
        import csv
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["tanggal", "nik", "nama", "bagian", "masuk", "pulang", "shift", "durasi", "status"])
        for r in data:
            w.writerow([r['tanggal'], r['nik'], r['nama'], r['bagian'],
                        r['masuk'], r['pulang'], r.get('shift', ''), _durasi(r['menit']), LABEL[r['status']]])
        nama = f"rekap_absensi_{dari}" + ("" if dari == sampai else f"_sd_{sampai}") + ".csv"
        return Response(buf.getvalue(), mimetype="text/csv",
                        headers={"Content-Disposition": f"attachment;filename={nama}"})
    if (request.args.get("format") or "xlsx").lower() == "pdf":
        # PDF ringkasan rapi (bukan template operator): 1 baris per orang per hari.
        baris = [[i + 1, r["tanggal"], r["nik"], r["nama"] or "-", r["bagian"] or "-",
                  r["masuk"] or "-", r["pulang"] or "-", r.get("shift", "") or "-",
                  _durasi(r["menit"]) or "-", r["jml"], LABEL[r["status"]]] for i, r in enumerate(data)]
        kolom = ["No", "Tanggal", "NIK", "Nama", "Bagian", "Masuk", "Pulang", "Shift",
                 "Durasi", "Scan", "Status"]
        nama = f"rekap_absensi_{dari}" + ("" if dari == sampai else f"_sd_{sampai}") + ".pdf"
        periode = dari.isoformat() if dari == sampai else f"{dari.isoformat()} s/d {sampai.isoformat()}"
        return _tabel_pdf("REKAP ABSENSI", "Periode: %s" % periode, kolom, baris, nama)
    # XLSX: isi template ABSENSI OPERATOR (format & rumus tidak diubah).
    # Bila template tidak ada, pakai format buatan lama sebagai cadangan.
    try:
        from laporan_template import isi_template, template_ada
        pakai_template = template_ada()
    except Exception:
        pakai_template = False
    if pakai_template:
        sendiri = None
        if session.get("role") != "admin":
            sendiri = (session.get("nik") or "").strip()
        try:
            buf, sheet, n_sel, n_orang, info = isi_template(dari, sampai, q, sendiri)
        except ValueError as e:
            return jsonify({"ok": False, "pesan": str(e)}), 400
        nama = (f"absensi_operator_{sampai:%Y-%m}_{sheet}"
                + ("" if dari == sampai else f"_{dari}_sd_{sampai}")) + ".xlsx"
        tak = info.get("tak_cocok", [])
        resp = Response(buf.getvalue(),
                        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f"attachment;filename={nama}"})
        # Info pencocokan agar web bisa tampilkan peringatan:
        # cocok = baris template berisi data mesin, tak_cocok = nama tanpa
        # data mesin (ditandai A, betulkan nama di tab Daftar).
        resp.headers["X-Absensi-Cocok"] = str(n_orang)
        resp.headers["X-Absensi-Sel"] = str(n_sel)
        resp.headers["X-Absensi-Tak-Cocok-N"] = str(len(tak))
        resp.headers["X-Absensi-Tak-Cocok"] = ("; ".join(tak[:20]))[:900]
        return resp
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    c = db()
    if session.get("role") != "admin":
        sendiri = (session.get("nik") or "").strip()
        detail = c.execute(
            """SELECT a.waktu, a.nik, COALESCE(p.nama,'') nama, COALESCE(p.bagian,'') bagian, a.tipe, a.cara
               FROM absen a LEFT JOIN pegawai p ON p.nik=a.nik
               WHERE a.nik=? AND date(a.waktu) BETWEEN ? AND ?
                 AND COALESCE(a.status,'DISETUJUI')='DISETUJUI'
               ORDER BY a.waktu""",
            (sendiri, dari.isoformat(), sampai.isoformat())).fetchall()
    else:
        detail = c.execute(
            """SELECT a.waktu, a.nik, COALESCE(p.nama,'') nama, COALESCE(p.bagian,'') bagian, a.tipe, a.cara
               FROM absen a LEFT JOIN pegawai p ON p.nik=a.nik
               WHERE date(a.waktu) BETWEEN ? AND ?
                 AND COALESCE(a.status,'DISETUJUI')='DISETUJUI'
               ORDER BY a.waktu""",
            (dari.isoformat(), sampai.isoformat())).fetchall()
    c.close()
    tinta = PatternFill("solid", fgColor="12343B")
    zebra = PatternFill("solid", fgColor="EEF4F3")
    kuning = PatternFill("solid", fgColor="FDF0D5")
    garis = Side(style="thin", color="C9D3D2")
    kotak = Border(left=garis, right=garis, top=garis, bottom=garis)

    def kepala(ws, baris, kolom):
        for i in range(1, kolom + 1):
            sel = ws.cell(row=baris, column=i)
            sel.fill, sel.font, sel.border = tinta, Font(color="FFFFFF", bold=True), kotak
            sel.alignment = Alignment(horizontal="center", vertical="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    periode = dari.isoformat() if dari == sampai else f"{dari.isoformat()} s/d {sampai.isoformat()}"
    ws.merge_cells("A1:K1")
    ws["A1"] = "REKAP ABSENSI"
    ws["A1"].font = Font(bold=True, size=14, color="12343B")
    ws.merge_cells("A2:K2")
    ws["A2"] = f"Periode: {periode}   |   Dicetak: {wib_now():%Y-%m-%d %H:%M}   |   Total scan: {len(detail)}"
    ws.append([])
    ws.append(["No", "Tanggal", "NIK", "Nama", "Bagian", "Jam Masuk", "Jam Pulang", "Shift",
               "Durasi", "Jml Scan", "Status"])
    kepala(ws, 4, 11)
    for no, r in enumerate(data, 1):
        ws.append([no, r["tanggal"], r["nik"], r["nama"] or "-", r["bagian"] or "-",
                   r["masuk"] or "-", r["pulang"] or "-", r.get("shift", "") or "-",
                   _durasi(r["menit"]) or "-", r["jml"], LABEL[r["status"]]])
        for i in range(1, 12):
            sel = ws.cell(row=ws.max_row, column=i)
            sel.border = kotak
            sel.alignment = Alignment(horizontal="left" if i == 4 else "center")
            if r["status"] in ("telat", "satu", "belum"):
                sel.fill = kuning
            elif no % 2 == 0:
                sel.fill = zebra
    if not data:
        ws.append(["-", "-", "-", "Belum ada data", "-", "-", "-", "-", "-", 0, "-"])
    for i, w in enumerate([6, 12, 12, 26, 16, 12, 12, 8, 10, 10, 15], 1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws2 = wb.create_sheet("Detail Scan")
    ws2.append(["Waktu", "NIK", "Nama", "Bagian", "Tipe", "Shift", "Sumber"])
    kepala(ws2, 1, 7)
    for r in detail:
        ws2.append([r["waktu"], r["nik"], r["nama"] or "-", r["bagian"] or "-",
                    r["tipe"], tentukan_shift(str(r["waktu"])[11:16]), r["cara"]])
        for i in range(1, 8):
            ws2.cell(row=ws2.max_row, column=i).border = kotak
    for i, w in enumerate([20, 12, 26, 16, 10, 8, 16], 1):
        ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:G{ws2.max_row}"

    buf = io.BytesIO()
    wb.save(buf)
    nama = f"rekap_absensi_{dari}" + ("" if dari == sampai else f"_sd_{sampai}") + ".xlsx"
    return Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment;filename={nama}"})


@bp_laporan.route("/backup")
@wajib_admin
def backup():
    if pakai_mysql():
        return Response(dump_sql(), mimetype="text/sql",
                        headers={"Content-Disposition":
                                 f"attachment;filename=absensi_backup_{wib_now():%Y%m%d_%H%M}.sql"})
    return send_file(DB, as_attachment=True,
                     download_name=f"absensi_backup_{wib_now():%Y%m%d_%H%M}.db")


def _tabel_xlsx(judul, subjudul, kolom, baris, nama_file):
    """Buat workbook rapi (gaya sama seperti export rekap): header gelap,
    border tipis, zebra, filter, freeze, judul + periode."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan"
    ws.merge_cells("A1:%s1" % chr(64 + len(kolom)))
    ws["A1"] = judul
    ws["A1"].font = Font(bold=True, size=14, color="12343B")
    ws.merge_cells("A2:%s2" % chr(64 + len(kolom)))
    ws["A2"] = "%s   |   Dicetak: %s   |   Total: %d" % (
        subjudul, wib_now().strftime("%Y-%m-%d %H:%M"), len(baris))
    ws.append([])
    ws.append(kolom)
    tinta = PatternFill("solid", fgColor="12343B")
    zebra = PatternFill("solid", fgColor="EEF4F3")
    garis = Side(style="thin", color="C9D3D2")
    kotak = Border(left=garis, right=garis, top=garis, bottom=garis)
    for i in range(1, len(kolom) + 1):
        sel = ws.cell(row=4, column=i)
        sel.fill, sel.font, sel.border = tinta, Font(color="FFFFFF", bold=True), kotak
        sel.alignment = Alignment(horizontal="center", vertical="center")
    for no, r in enumerate(baris, 1):
        ws.append(r)
        for i in range(1, len(kolom) + 1):
            sel = ws.cell(row=ws.max_row, column=i)
            sel.border = kotak
            sel.alignment = Alignment(horizontal="center", vertical="center")
            if no % 2 == 0:
                sel.fill = zebra
    if not baris:
        ws.append(["-"] * len(kolom))
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:%s%d" % (chr(64 + len(kolom)), ws.max_row)
    # lebar kolom otomatis dari isi
    for i in range(1, len(kolom) + 1):
        mx = max([len(str(kolom[i - 1]))] + [len(str(ws.cell(r, i).value or ""))
                                             for r in range(4, ws.max_row + 1)])
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = min(42, max(10, mx + 2))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment;filename=%s" % nama_file})


def _tabel_csv(kolom, baris, nama_file):
    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(kolom)
    w.writerows(baris)
    return Response(buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment;filename=%s" % nama_file})


def _tabel_pdf(judul, subjudul, kolom, baris, nama_file):
    """Laporan PDF rapi (landscape A4): judul + subjudul, header gelap,
    zebra, garis, nomor halaman, header diulang tiap halaman."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                     Paragraph)
    from reportlab.lib.enums import TA_CENTER
    GELAP = colors.HexColor("#12343B")
    ZEBRA = colors.HexColor("#EEF4F3")
    GARIS = colors.HexColor("#C9D3D2")
    ABU = colors.HexColor("#5B6B6B")

    halaman = landscape(A4)
    margin = 12 * mm
    lebar_pakai = halaman[0] - 2 * margin
    # lebar kolom proporsional isi (min 40, maks 170)
    pj = []
    for i, h in enumerate(kolom):
        mx = len(str(h))
        for r in baris[:200]:
            try:
                mx = max(mx, len(str(r[i])))
            except IndexError:
                pass
        pj.append(mx)
    tot = sum(pj) or 1
    lebar = [max(40, min(170, lebar_pakai * p / tot)) for p in pj]
    selisih = lebar_pakai - sum(lebar)
    if abs(selisih) > 1:
        lebar[-1] += selisih

    g_judul = ParagraphStyle("judul", fontName="Helvetica-Bold", fontSize=15,
                             textColor=GELAP, spaceAfter=2)
    g_sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=9,
                           textColor=ABU, spaceAfter=8)
    g_head = ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=8,
                            textColor=colors.white, alignment=TA_CENTER)
    g_sel = ParagraphStyle("sel", fontName="Helvetica", fontSize=8,
                           leading=10, alignment=TA_CENTER)

    def P(teks, g):
        return Paragraph(str(teks if teks != "" else "-").replace("&", "&amp;")
                         .replace("<", "&lt;").replace(">", "&gt;"), g)

    isi = [[P(h, g_head) for h in kolom]]
    if baris:
        for r in baris:
            isi.append([P(r[i] if i < len(r) else "", g_sel) for i in range(len(kolom))])
    else:
        isi.append([P("Belum ada data pada periode ini.", g_sel)] + [P("", g_sel)] * (len(kolom) - 1))

    gaya = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), GELAP),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, GARIS),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ])
    for i in range(1, len(isi)):
        if i % 2 == 0:
            gaya.add("BACKGROUND", (0, i), (-1, i), ZEBRA)
    tabel = Table(isi, colWidths=lebar, repeatRows=1)
    tabel.setStyle(gaya)

    buf = io.BytesIO()
    dok = SimpleDocTemplate(buf, pagesize=halaman,
                            leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=15 * mm,
                            title=judul, author="Absensi Online IPP")

    def kaki(kanvas, _dok):
        kanvas.saveState()
        kanvas.setFont("Helvetica", 8)
        kanvas.setFillColor(ABU)
        kanvas.drawCentredString(halaman[0] / 2, 10 * mm, "Halaman %d" % _dok.page)
        kanvas.restoreState()

    dok.build([Paragraph(judul, g_judul),
               Paragraph(subjudul + "   |   Dicetak: %s   |   Total: %d baris"
                         % (wib_now().strftime("%Y-%m-%d %H:%M"), len(baris)), g_sub),
               tabel], onFirstPage=kaki, onLaterPages=kaki)
    buf.seek(0)
    return Response(buf.getvalue(), mimetype="application/pdf",
                    headers={"Content-Disposition": "attachment;filename=%s" % nama_file})


@bp_laporan.route("/export/hari")
@wajib_admin
def export_hari():
    """Unduh laporan HARI INI / 1 tanggal: xlsx + csv rapi."""
    from config import JAM_MASUK
    tgl = (request.args.get("tgl") or request.args.get("dari")
           or wib_today().isoformat())[:10]
    c = db()
    peg = [dict(r) for r in c.execute(
        "SELECT nik, nama, COALESCE(bagian,'') bagian FROM pegawai ORDER BY COALESCE(nama,nik)").fetchall()]
    absen = c.execute(
        "SELECT nik, MIN(waktu) masuk, MAX(waktu) pulang, COUNT(*) n "
        "FROM absen WHERE date(waktu)=date(?) AND COALESCE(status,'DISETUJUI')='DISETUJUI' "
        "GROUP BY nik", (tgl,)).fetchall()
    a = {r["nik"]: dict(r) for r in absen}
    izin = {r["nik"]: r["jenis"] for r in c.execute(
        "SELECT nik,jenis FROM pengajuan WHERE tanggal=? AND status='DISETUJUI'",
        (tgl,)).fetchall()}
    c.close()
    baris = []
    for no, p in enumerate(peg, 1):
        nik = p["nik"]
        if nik in izin:
            baris.append([no, nik, p["nama"], p["bagian"], "-", "-", "-", izin[nik]])
        elif nik not in a:
            baris.append([no, nik, p["nama"], p["bagian"], "-", "-", "-", "Alpha"])
        else:
            m = a[nik]["masuk"][11:16]
            pl = a[nik]["pulang"][11:16] if a[nik]["n"] > 1 else "-"
            st = "Hadir" if m <= JAM_MASUK else "Telat %s" % m
            if a[nik]["n"] < 2:
                st += " (1x tap)"
            baris.append([no, nik, p["nama"], p["bagian"], m, pl, tentukan_shift(m), st])
    kolom = ["No", "NIK", "Nama", "Bagian", "Jam Masuk", "Jam Pulang", "Shift", "Status"]
    if (request.args.get("format") or "xlsx").lower() == "csv":
        return _tabel_csv(kolom, baris, "laporan_harian_%s.csv" % tgl)
    if (request.args.get("format") or "").lower() == "pdf":
        return _tabel_pdf("LAPORAN HARIAN", "Tanggal: %s" % tgl, kolom, baris,
                          "laporan_harian_%s.pdf" % tgl)
    return _tabel_xlsx("LAPORAN HARIAN", "Tanggal: %s" % tgl, kolom, baris,
                       "laporan_harian_%s.xlsx" % tgl)


@bp_laporan.route("/export/izin")
@wajib_admin
def export_izin():
    """Unduh daftar IZIN / pengajuan: xlsx + csv rapi."""
    st = (request.args.get("status") or "").upper()
    c = db()
    q = ("SELECT p.nik, COALESCE(g.nama,p.nik) nama, COALESCE(g.bagian,'') bagian, "
         "p.tanggal, p.jenis, COALESCE(p.ket,'') ket, p.status "
         "FROM pengajuan p LEFT JOIN pegawai g ON g.nik=p.nik")
    args = []
    if st in ("PENDING", "DISETUJUI", "DITOLAK"):
        q += " WHERE p.status=?"
        args.append(st)
    q += " ORDER BY p.tanggal DESC, p.id DESC"
    rows = c.execute(q, args).fetchall()
    c.close()
    baris = [[i + 1, r["nik"], r["nama"], r["bagian"], r["tanggal"],
              r["jenis"], r["ket"] or "-", r["status"]] for i, r in enumerate(rows)]
    kolom = ["No", "NIK", "Nama", "Bagian", "Tanggal", "Jenis", "Alasan", "Status"]
    nama = "laporan_izin_%s" % (st.lower() if st else "semua")
    if (request.args.get("format") or "xlsx").lower() == "csv":
        return _tabel_csv(kolom, baris, nama + ".csv")
    if (request.args.get("format") or "").lower() == "pdf":
        return _tabel_pdf("LAPORAN IZIN", "Status: %s" % (st or "SEMUA"),
                          kolom, baris, nama + ".pdf")
    return _tabel_xlsx("LAPORAN IZIN", "Status: %s" % (st or "SEMUA"),
                       kolom, baris, nama + ".xlsx")


@bp_laporan.route("/export/pegawai")
@wajib_admin
def export_pegawai():
    """Unduh DAFTAR pegawai: xlsx + csv rapi."""
    q = (request.args.get("q") or "").strip()
    like = "%" + q + "%"
    c = db()
    rows = c.execute(
        "SELECT nik, nama, COALESCE(bagian,'') bagian FROM pegawai "
        "WHERE ?='' OR nik LIKE ? OR nama LIKE ? OR bagian LIKE ? ORDER BY nik",
        (q, like, like, like)).fetchall()
    # hitung kehadiran per orang agar sheet berguna (hanya DISETUJUI)
    stat = {r["nik"]: (r["tap"], r["hari"]) for r in c.execute(
        "SELECT nik, COUNT(*) tap, COUNT(DISTINCT substr(waktu,1,10)) hari "
        "FROM absen WHERE COALESCE(status,'DISETUJUI')='DISETUJUI' GROUP BY nik").fetchall()}
    c.close()
    baris = [[i + 1, r["nik"], r["nama"], r["bagian"],
              stat.get(r["nik"], (0, 0))[0], stat.get(r["nik"], (0, 0))[1]]
             for i, r in enumerate(rows)]
    kolom = ["No", "NIK", "Nama", "Bagian", "Total Tap", "Total Hari Hadir"]
    if (request.args.get("format") or "xlsx").lower() == "csv":
        return _tabel_csv(kolom, baris, "daftar_pegawai.csv")
    if (request.args.get("format") or "").lower() == "pdf":
        return _tabel_pdf("DAFTAR PEGAWAI", "Total: %d orang" % len(baris),
                          kolom, baris, "daftar_pegawai.pdf")
    return _tabel_xlsx("DAFTAR PEGAWAI", "Total: %d orang" % len(baris),
                       kolom, baris, "daftar_pegawai.xlsx")


def _rentang_hapus(scope, d):
    if scope == "harian":
        return d, d
    if scope == "mingguan":
        awal = d - timedelta(days=d.weekday())
        return awal, awal + timedelta(days=6)
    if scope == "bulanan":
        return d.replace(day=1), d.replace(day=monthrange(d.year, d.month)[1])
    if scope == "tahunan":
        return date(d.year, 1, 1), date(d.year, 12, 31)
    return None


@bp_laporan.route("/api/reset", methods=["POST"])
@wajib_admin
def api_reset():
    j = request.get_json(force=True, silent=True) or {}
    r = _rentang_hapus(j.get("scope", "harian"), _tanggal(j.get("date"), wib_today()))
    if not r:
        return jsonify(ok=False, pesan="Jenis reset tidak dikenal.")
    a, b = r[0].isoformat(), r[1].isoformat()
    teks = a if a == b else f"{a} s/d {b}"
    c = db()
    n = c.execute("SELECT COUNT(*) FROM absen WHERE date(waktu) BETWEEN ? AND ?",
                  (a, b)).fetchone()[0]
    if j.get("dry", False):  # hanya hitung dulu, belum menghapus
        c.close()
        return jsonify(ok=True, jumlah=n, rentang=teks)
    c.close()
    os.makedirs(BACKUP_DIR, exist_ok=True)
    cadangan = os.path.join(BACKUP_DIR,
                            f"absensi_sebelum_hapus_{wib_now():%Y%m%d_%H%M%S}" + (".sql" if pakai_mysql() else ".db"))
    snapshot_db(cadangan)
    c = db()
    c.execute("DELETE FROM absen WHERE date(waktu) BETWEEN ? AND ?", (a, b))
    c.commit()
    c.close()
    return jsonify(ok=True, hapus=n,
                   pesan=f"{n} scan ({teks}) dihapus. Backup tersimpan di folder backup.")
